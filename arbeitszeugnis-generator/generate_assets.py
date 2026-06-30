#!/usr/bin/env python3
"""Erzeugt DEMO-Assets (Satzdatenbank + Word-Vorlagen), damit die App ohne die
echten, proprietären Dateien lauffähig ist.

WICHTIG – Schutz echter Daten:
  * Vorhandene Dateien werden NIE überschrieben. Liegt deine echte
    ``data/satzdatenbank.xlsx`` oder eine echte WKVZ-Vorlage bereits da, wird
    sie übersprungen.
  * Die Demo-Sätze sind frei erfunden und nur als Strukturbeispiel gedacht.
    Für den Echtbetrieb deine eigene Satzdatenbank und die WKVZ-Vorlagen ablegen.

Demo-Excel-Struktur (entspricht dem Echtformat):
  Blätter: Führungskräfte | Mitarbeitende | Lehrpersonen
  Spalten: Kompetenzbereich | Beobachtungsmerkmal | 100 | 200 | 300 | 400
  (100 = bestes; zusammengeführte Kompetenzbereich-Zellen bleiben leer ->
   werden beim Einlesen per forward-fill aufgefüllt.)
"""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from docx import Document
from docx.shared import Pt

# Pfad-Setup, damit das Skript auch per Doppelklick / aus beliebigem CWD läuft.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app import config  # noqa: E402


# ---------------------------------------------------------------------------
# Demo-Satzdaten (frei erfunden, nur Strukturbeispiel)
# ---------------------------------------------------------------------------
# Pro Merkmal: (Note 100, 200, 300, 400). Einzelne Sätze haben bewusst KEINEN
# Schlusspunkt, um die punktuationssichere Verbindung zu testen.
LEHRPERSON = {
    "Unterrichtsqualität": {
        "Fachkompetenz": (
            "[Name] verfügt über ein ausgezeichnetes, stets aktuelles Fachwissen, das sie/er didaktisch überzeugend in ihren/seinen Unterricht einbringt.",
            "[Name] verfügt über ein fundiertes Fachwissen, das sie/er sicher in ihren/seinen Unterricht einbringt.",
            "[Name] verfügt über solides Fachwissen und setzt es im Unterricht sachgerecht ein.",
            "[Name] verfügt über grundlegendes Fachwissen.",
        ),
        "Methodik und Didaktik": (
            "Ihren/Seinen Unterricht gestaltet sie/er methodisch vielfältig und stets adressatengerecht",  # ohne Punkt (Absicht)
            "Ihren/Seinen Unterricht gestaltet sie/er methodisch abwechslungsreich und meist adressatengerecht.",
            "[Name] setzt verschiedene Unterrichtsmethoden zweckmässig ein.",
            "[Name] setzt einzelne Unterrichtsmethoden ein.",
        ),
        "Unterrichtsgestaltung": (
            "[Name] schafft eine anregende Lernatmosphäre und begeistert die Lernenden für ihre/seine Fächer.",
            "[Name] schafft eine gute Lernatmosphäre und motiviert die Lernenden.",
            "[Name] sorgt für eine ruhige Lernatmosphäre.",
            "[Name] bemüht sich um eine geordnete Lernatmosphäre.",
        ),
        "Leistungsbeurteilung": (
            "Leistungen beurteilt sie/er transparent, fair und nachvollziehbar.",
            "Leistungen beurteilt sie/er nachvollziehbar und fair.",
            "[Name] beurteilt Leistungen korrekt.",
            "[Name] beurteilt Leistungen nach Vorgabe.",
        ),
    },
    "Schülerorientierung": {
        "Beziehung zu den Lernenden": (
            "Zu den Lernenden pflegt sie/er ein wertschätzendes, von Vertrauen geprägtes Verhältnis.",
            "Zu den Lernenden pflegt sie/er ein gutes, respektvolles Verhältnis.",
            "[Name] pflegt ein korrektes Verhältnis zu den Lernenden.",
            "[Name] pflegt ein distanziertes Verhältnis zu den Lernenden.",
        ),
        "Förderung und Differenzierung": (
            "Lernende mit unterschiedlichen Voraussetzungen fördert sie/er gezielt und differenziert.",
            "Lernende fördert sie/er aufmerksam und individuell.",
            "[Name] fördert die Lernenden im Rahmen der Möglichkeiten.",
            "[Name] fördert die Lernenden nach Vorgabe.",
        ),
        "Klassenführung": (
            "Ihre/Seine Klasse führt sie/er umsichtig und mit natürlicher Autorität.",
            "Ihre/Seine Klasse führt sie/er sicher und konsequent",  # ohne Punkt (Absicht)
            "[Name] führt die Klasse mehrheitlich sicher.",
            "[Name] führt die Klasse mit Unterstützung.",
        ),
    },
    "Zusammenarbeit und Engagement": {
        "Zusammenarbeit im Kollegium": (
            "Im Kollegium arbeitet sie/er ausgesprochen kooperativ und ist eine geschätzte Ansprechperson.",
            "Im Kollegium arbeitet sie/er kooperativ und zuverlässig.",
            "[Name] arbeitet im Kollegium korrekt mit.",
            "[Name] arbeitet im Kollegium nach Aufforderung mit.",
        ),
        "Elternarbeit": (
            "Die Zusammenarbeit mit den Eltern gestaltet sie/er professionell und vertrauensbildend.",
            "Die Zusammenarbeit mit den Eltern gestaltet sie/er sachlich und verlässlich.",
            "[Name] pflegt einen korrekten Kontakt mit den Eltern.",
            "[Name] pflegt den Kontakt mit den Eltern nach Bedarf.",
        ),
        "Engagement und Schulentwicklung": (
            "Über den Unterricht hinaus engagiert sie/er sich mit grossem Einsatz für die Schulentwicklung.",
            "Über den Unterricht hinaus engagiert sie/er sich für gemeinsame Anliegen der Schule.",
            "[Name] beteiligt sich an schulischen Anlässen.",
            "[Name] beteiligt sich nach Aufforderung an schulischen Anlässen.",
        ),
    },
    "Organisation und Professionalität": {
        "Zuverlässigkeit und Organisation": (
            "Ihre/Seine Arbeit erledigt sie/er äusserst zuverlässig, sorgfältig und gut organisiert.",
            "Ihre/Seine Arbeit erledigt sie/er zuverlässig und sorgfältig.",
            "[Name] erledigt die Arbeit zuverlässig.",
            "[Name] erledigt die Arbeit mit Erinnerung zuverlässig.",
        ),
        "Reflexion und Weiterbildung": (
            "[Name] reflektiert ihr/sein Handeln kritisch und bildet sich kontinuierlich weiter.",
            "[Name] reflektiert ihre/seine Arbeit und bildet sich regelmässig weiter.",
            "[Name] nimmt an Weiterbildungen teil.",
            "[Name] nimmt an obligatorischen Weiterbildungen teil.",
        ),
        "Auftreten und Vorbildfunktion": (
            "[Name] tritt jederzeit professionell auf und ist den Lernenden ein glaubwürdiges Vorbild.",
            "[Name] tritt professionell auf und verhält sich vorbildlich.",
            "[Name] tritt korrekt auf.",
            "[Name] tritt zurückhaltend auf.",
        ),
    },
}

# Minimal-Demo für die noch nicht ausgebauten Zielgruppen (nur Struktur).
MITARBEITENDE = {
    "Arbeitsqualität und Leistung": {
        "Fachkompetenz": (
            "[Name] verfügt über ein ausgezeichnetes Fachwissen und setzt es souverän ein.",
            "[Name] verfügt über ein gutes Fachwissen und setzt es sicher ein.",
            "[Name] verfügt über solides Fachwissen.",
            "[Name] verfügt über grundlegendes Fachwissen.",
        ),
        "Arbeitsweise": (
            "[Name] arbeitet äusserst sorgfältig, speditiv und zuverlässig.",
            "[Name] arbeitet sorgfältig und zuverlässig.",
            "[Name] arbeitet zuverlässig.",
            "[Name] arbeitet nach Anleitung zuverlässig.",
        ),
    },
    "Verhalten": {
        "Zusammenarbeit": (
            "Im Team arbeitet sie/er ausgesprochen kooperativ und ist sehr geschätzt.",
            "Im Team arbeitet sie/er kooperativ und verlässlich.",
            "[Name] arbeitet im Team korrekt mit.",
            "[Name] arbeitet im Team nach Aufforderung mit.",
        ),
    },
}

FUEHRUNGSKRAFT = {
    "Führung": {
        "Mitarbeiterführung": (
            "[Name] führt ihr/sein Team motivierend, klar und mit grossem Vertrauen.",
            "[Name] führt ihr/sein Team verlässlich und zielorientiert.",
            "[Name] führt das Team korrekt.",
            "[Name] führt das Team mit Unterstützung.",
        ),
        "Entscheidungsverhalten": (
            "Entscheidungen trifft sie/er umsichtig, zeitgerecht und nachvollziehbar.",
            "Entscheidungen trifft sie/er überlegt und zeitgerecht.",
            "[Name] trifft Entscheidungen sachgerecht.",
            "[Name] trifft Entscheidungen nach Rücksprache.",
        ),
    },
    "Strategie und Ergebnis": {
        "Zielerreichung": (
            "Vereinbarte Ziele erreicht sie/er regelmässig und übertrifft sie häufig.",
            "Vereinbarte Ziele erreicht sie/er zuverlässig.",
            "[Name] erreicht die vereinbarten Ziele grösstenteils.",
            "[Name] erreicht die vereinbarten Ziele teilweise.",
        ),
    },
}


def _schreibe_blatt(ws, daten: dict) -> None:
    kopf = ["Kompetenzbereich", "Beobachtungsmerkmal"] + [str(c) for c in config.GRADE_CODES]
    ws.append(kopf)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")

    for bereich, merkmale in daten.items():
        erste_zeile_des_bereichs = True
        for merkmal, saetze in merkmale.items():
            # Kompetenzbereich nur in der ersten Zeile der Gruppe -> testet ffill.
            kb = bereich if erste_zeile_des_bereichs else None
            ws.append([kb, merkmal, *saetze])
            erste_zeile_des_bereichs = False

    # Spaltenbreiten grosszügig, damit die Datei lesbar bleibt.
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30
    for col in ("C", "D", "E", "F"):
        ws.column_dimensions[col].width = 60


def erzeuge_satzdatenbank(pfad: str) -> bool:
    if os.path.exists(pfad):
        print(f"  übersprungen (vorhanden): {os.path.relpath(pfad, config.BASE_DIR)}")
        return False
    wb = Workbook()
    wb.remove(wb.active)
    _schreibe_blatt(wb.create_sheet("Führungskräfte"), FUEHRUNGSKRAFT)
    _schreibe_blatt(wb.create_sheet("Mitarbeitende"), MITARBEITENDE)
    _schreibe_blatt(wb.create_sheet("Lehrpersonen"), LEHRPERSON)
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    wb.save(pfad)
    print(f"  erstellt: {os.path.relpath(pfad, config.BASE_DIR)}")
    return True


# ---------------------------------------------------------------------------
# Demo-Word-Vorlagen
# ---------------------------------------------------------------------------
def _absatz(doc, text, bold=False, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def erzeuge_lehrperson_vorlage(pfad: str) -> bool:
    if os.path.exists(pfad):
        print(f"  übersprungen (vorhanden): {os.path.relpath(pfad, config.BASE_DIR)}")
        return False
    doc = Document()
    _absatz(doc, "DEMO-Vorlage – durch echte WKVZ-Vorlage ersetzen", bold=True, size=9)
    _absatz(doc, "Arbeitszeugnis", bold=True, size=16)
    doc.add_paragraph()
    _absatz(doc, "{{ anrede }} {{ voller_name }}")
    _absatz(doc, "{{ strasse_hausnummer }}")
    _absatz(doc, "{{ plz_ort }}")
    doc.add_paragraph()
    _absatz(
        doc,
        "{{ anrede }} {{ voller_name }}, geboren am {{ geburtsdatum }}, von "
        "{{ heimatort }}, war vom {{ eintrittsdatum }} bis {{ austrittsdatum }} "
        "als {{ funktion }} (Pensum {{ pensum }}) bei uns tätig. "
        "Unterrichtet wurden die Fachbereiche {{ fachbereiche }} in den Klassen {{ klassen }}.",
    )
    _absatz(doc, "{{ vertrag_absatz }}")
    doc.add_paragraph()
    _absatz(doc, "Unterrichtsqualität", bold=True)
    _absatz(doc, "{{ absatz_unterrichtsqualitaet }}")
    doc.add_paragraph()
    _absatz(doc, "Sozialverhalten und Zusammenarbeit", bold=True)
    _absatz(doc, "{{ absatz_sozialverhalten }}")
    doc.add_paragraph()
    _absatz(doc, "{{ zusatz_absatz }}")
    _absatz(doc, "{{ schlusssatz }}")
    doc.add_paragraph()
    _absatz(doc, "Ort, {{ ausstellungsdatum }}")
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    doc.save(pfad)
    print(f"  erstellt: {os.path.relpath(pfad, config.BASE_DIR)}")
    return True


def erzeuge_platzhalter_vorlage(pfad: str, titel: str) -> bool:
    """Knappe Demo-Vorlage für FK/MA (noch nicht ausgebaut)."""
    if os.path.exists(pfad):
        print(f"  übersprungen (vorhanden): {os.path.relpath(pfad, config.BASE_DIR)}")
        return False
    doc = Document()
    _absatz(doc, f"DEMO-Vorlage ({titel}) – noch nicht ausgebaut", bold=True, size=9)
    _absatz(doc, "Arbeitszeugnis", bold=True, size=16)
    doc.add_paragraph()
    _absatz(doc, "{{ anrede }} {{ voller_name }}")
    _absatz(doc, "{{ funktion }}, {{ eintrittsdatum }} – {{ austrittsdatum }}")
    _absatz(doc, "{{ schlusssatz }}")
    os.makedirs(os.path.dirname(pfad), exist_ok=True)
    doc.save(pfad)
    print(f"  erstellt: {os.path.relpath(pfad, config.BASE_DIR)}")
    return True


def main() -> None:
    print("Demo-Assets erzeugen (echte Dateien werden nicht überschrieben):")
    erzeuge_satzdatenbank(config.SATZDATENBANK)
    erzeuge_lehrperson_vorlage(config.TEMPLATES["lehrperson"])
    erzeuge_platzhalter_vorlage(config.TEMPLATES["mitarbeitende"], "Mitarbeitende")
    erzeuge_platzhalter_vorlage(config.TEMPLATES["fuehrungskraft"], "Führungskräfte")
    print("Fertig.")


if __name__ == "__main__":
    main()
